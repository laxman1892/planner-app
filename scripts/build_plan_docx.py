from __future__ import annotations

import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


TITLE_COLOR = RGBColor(0x0B, 0x25, 0x45)
H1_COLOR = RGBColor(0x2E, 0x74, 0xB5)
H2_COLOR = RGBColor(0x2E, 0x74, 0xB5)
H3_COLOR = RGBColor(0x1F, 0x4D, 0x78)
MUTED_COLOR = RGBColor(0x55, 0x55, 0x55)


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def configure_page(section) -> None:
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.49)
    section.footer_distance = Inches(0.49)


def ensure_styles(doc: Document) -> None:
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.1

    title = doc.styles["Title"]
    title.font.name = "Calibri"
    title.font.size = Pt(22)
    title.font.bold = True
    title.font.color.rgb = TITLE_COLOR
    title.paragraph_format.space_after = Pt(8)

    subtitle = doc.styles["Subtitle"]
    subtitle.font.name = "Calibri"
    subtitle.font.size = Pt(11)
    subtitle.font.color.rgb = MUTED_COLOR

    for style_name, size, color in (
        ("Heading 1", 16, H1_COLOR),
        ("Heading 2", 13, H2_COLOR),
        ("Heading 3", 12, H3_COLOR),
    ):
        style = doc.styles[style_name]
        style.font.name = "Calibri"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = color
        style.paragraph_format.space_before = Pt(12)
        style.paragraph_format.space_after = Pt(6)

    if "Phase Summary" not in doc.styles:
        style = doc.styles.add_style("Phase Summary", WD_STYLE_TYPE.PARAGRAPH)
        style.base_style = doc.styles["Normal"]
        style.font.name = "Calibri"
        style.font.size = Pt(10.5)
        style.paragraph_format.space_after = Pt(4)


def add_header(section, text: str) -> None:
    paragraph = section.header.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run(text)
    run.font.name = "Calibri"
    run.font.size = Pt(9)
    run.font.color.rgb = MUTED_COLOR


def add_cover(doc: Document) -> None:
    p = doc.add_paragraph(style="Title")
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.add_run("PlanQuest Developer Execution Plan")

    subtitle = doc.add_paragraph(style="Subtitle")
    subtitle.add_run(
        "Detailed implementation sequence for upgrading the current project into a production-style collaborative planner portfolio app."
    )

    meta = doc.add_paragraph()
    meta.style = doc.styles["Phase Summary"]
    meta.add_run("Prepared: 2026-05-29\n")
    meta.add_run("Audience: solo developer / future collaborators\n")
    meta.add_run("Source of truth: docs/superpowers/plans/2026-05-29-planquest-master-execution-plan.md")

    doc.add_paragraph("")

    table = doc.add_table(rows=1, cols=4)
    table.style = "Table Grid"
    table.autofit = False
    widths = [Inches(1.05), Inches(2.3), Inches(1.35), Inches(1.8)]
    headers = ["Phase", "Primary Outcome", "Workload", "Exit Signal"]
    header = table.rows[0]
    set_repeat_table_header(header)
    for cell, text, width in zip(header.cells, headers, widths):
        cell.width = width
        cell.text = text
        set_cell_shading(cell, "E8EEF5")
        for paragraph in cell.paragraphs:
            for run in paragraph.runs:
                run.font.bold = True
                run.font.size = Pt(10)
    rows = [
        ("1", "Lock product direction", "Low code / high clarity", "MVP boundaries approved"),
        ("2", "Refactor frontend architecture", "Heavy refactor", "Domain-based UI structure"),
        ("3", "Harden collaboration backend", "Heavy backend", "Two-user rules enforced"),
        ("4", "Add quality workflow", "Tooling focused", "CI + lint + test gates live"),
        ("5", "Ship event collaboration slice", "Full-stack delivery", "Two-user flow demo-ready"),
        ("6", "Rework challenges", "Feature normalization", "Challenge logic coherent"),
        ("7", "Prepare deployment", "Ops and release", "Portfolio-grade release"),
    ]
    for row_data in rows:
        row = table.add_row()
        for cell, text, width in zip(row.cells, row_data, widths):
            cell.width = width
            cell.text = text

    doc.add_page_break()


def add_section_heading(doc: Document, line: str) -> None:
    hashes = len(line) - len(line.lstrip("#"))
    text = line[hashes:].strip()
    if hashes == 1:
        doc.add_paragraph(text, style="Heading 1")
    elif hashes == 2:
        doc.add_paragraph(text, style="Heading 2")
    else:
        doc.add_paragraph(text, style="Heading 3")


def clean_inline(text: str) -> str:
    text = re.sub(r"`([^`]+)`", r"\1", text)
    text = re.sub(r"\*\*([^*]+)\*\*", r"\1", text)
    return text.strip()


def add_bullet(doc: Document, text: str, level: int = 0) -> None:
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.left_indent = Inches(0.25 + (0.2 * level))
    p.paragraph_format.space_after = Pt(4)
    p.add_run(clean_inline(text))


def add_number(doc: Document, text: str) -> None:
    p = doc.add_paragraph(style="List Number")
    p.paragraph_format.space_after = Pt(4)
    p.add_run(clean_inline(text))


def add_paragraph(doc: Document, text: str) -> None:
    p = doc.add_paragraph(style="Normal")
    p.add_run(clean_inline(text))


def render_markdown_into_doc(doc: Document, markdown_text: str) -> None:
    lines = markdown_text.splitlines()
    for raw_line in lines:
        line = raw_line.rstrip()
        if not line.strip():
            continue
        if line.startswith("> "):
            p = doc.add_paragraph(style="Phase Summary")
            run = p.add_run(clean_inline(line[2:]))
            run.italic = True
            continue
        if line.startswith("#"):
            add_section_heading(doc, line)
            continue
        if re.match(r"^\d+\.\s+", line):
            add_number(doc, re.sub(r"^\d+\.\s+", "", line))
            continue
        if line.startswith("- [ ] "):
            add_bullet(doc, line[6:])
            continue
        if line.startswith("- "):
            add_bullet(doc, line[2:])
            continue
        add_paragraph(doc, line)


def build_docx(markdown_path: Path, output_path: Path) -> None:
    doc = Document()
    configure_page(doc.sections[0])
    ensure_styles(doc)
    add_header(doc.sections[0], "PlanQuest Execution Plan")
    add_cover(doc)

    new_section = doc.add_section(WD_SECTION.NEW_PAGE)
    configure_page(new_section)
    add_header(new_section, "PlanQuest Execution Plan")

    render_markdown_into_doc(doc, markdown_path.read_text(encoding="utf-8"))
    doc.save(output_path)


def main(argv: list[str]) -> int:
    if len(argv) != 3:
        print("Usage: build_plan_docx.py <input.md> <output.docx>")
        return 1

    markdown_path = Path(argv[1]).resolve()
    output_path = Path(argv[2]).resolve()
    output_path.parent.mkdir(parents=True, exist_ok=True)
    build_docx(markdown_path, output_path)
    print(output_path)
    return 0


if __name__ == "__main__":
    raise SystemExit(main(sys.argv))
